"""Сравнение нашего каталога запчастей с прайсами поставщиков.

Запуск из корня репозитория:

    python compare.py

Настройки — config/settings.yaml (наш каталог, курсы, правила),
config/suppliers.yaml (как читать каждый прайс), config/brands.yaml
(единое написание производителей). Результат — result/price_comparison_<дата>.xlsx.
"""

from __future__ import annotations

import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from statistics import mean, median

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CONFIG = ROOT / "config"


# ---------------------------------------------------------------- конфиг


def load_yaml(name: str) -> dict:
    with open(CONFIG / name, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def resolve_file(pattern: str) -> Path:
    """Путь или шаблон со * → один файл (последний по имени)."""
    matches = sorted(p for p in ROOT.glob(pattern) if p.is_file() and not p.name.startswith("~$"))
    if not matches:
        sys.exit(f"Файл не найден: {pattern}")
    if len(matches) > 1:
        print(f"  по шаблону {pattern} несколько файлов, беру {matches[-1].name}")
    return matches[-1]


# ---------------------------------------------------------------- нормализация

# Кириллица, которая выглядит как латиница: в артикулах их путают.
LOOKALIKE = str.maketrans("АВЕКМНОРСТХУ", "ABEKMHOPCTXY")
ARTICLE_JUNK = re.compile(r"[\s\-./\\_]+")


def cell_text(v) -> str:
    """Значение ячейки как строка; целое число из Excel — без «.0»."""
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        if v.is_integer():
            return str(int(v))
    return str(v).replace("\xa0", " ").strip()


def norm_article(v) -> str:
    return ARTICLE_JUNK.sub("", cell_text(v).upper().translate(LOOKALIKE))


def parse_number(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    s = str(v).replace("\xa0", "").replace(" ", "").lower()
    s = re.sub(r"(руб\.?|р\.|₽|₸|тг\.?|kzt|rub)", "", s).replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def parse_price(v) -> float | None:
    p = parse_number(v)
    return p if p and p > 0 else None


# Заводской номер ГАЗ/УАЗ: модель-[модификация-]деталь[-исполнение],
# например 3151-00-3501068-10, 469-1701164, 24-1000102-31.
SEPARATED = re.compile(r"^(\d{2,4})[\s\-.]+(?:(\d{2})[\s\-.]+)?(\d{7})(?:[\s\-.]*(\d{2,3}))?$")


def factory_key(model: str, mod: str, detail: str, suffix: str) -> str:
    # Модификация 00 и исполнение 00 в записи часто опускают.
    return "|".join((model.lstrip("0"), "" if mod in ("", "00") else mod, detail, suffix.lstrip("0")))


def factory_keys(raw) -> set[str]:
    """Ключи «формата ГАЗ/УАЗ» для сопоставления записей с разными разделителями.

    С разделителями номер разбирается однозначно. Сплошные цифры
    (у нас артикулы хранятся числом) разбираются всеми возможными способами.
    """
    text = cell_text(raw).upper()
    m = SEPARATED.match(text)
    if m:
        model, mod, detail, suffix = (g or "" for g in m.groups())
        return {factory_key(model, mod, detail, suffix)}
    digits = norm_article(raw)
    if not digits.isdigit():
        return set()
    digits = digits.lstrip("0")
    keys = set()
    for ml in (2, 3, 4):
        for modl in (0, 2):
            for sl in (0, 2, 3):
                if ml + modl + 7 + sl != len(digits):
                    continue
                model = digits[:ml]
                mod = digits[ml : ml + modl]
                detail = digits[ml + modl : ml + modl + 7]
                suffix = digits[ml + modl + 7 :]
                keys.add(factory_key(model, mod, detail, suffix))
    return keys


def brand_text(v) -> str:
    s = cell_text(v).upper().replace("«", '"').replace("»", '"').replace('"', "")
    return re.sub(r"\s+", " ", s).strip()


class Brands:
    """Единое написание производителей и их тип: оригинал / аналог / не бренд."""

    ORIGINAL, ANALOG, UNKNOWN = "оригинал", "аналог", "не указан"

    def __init__(self, cfg: dict):
        self.alias: dict[str, tuple[str, bool]] = {}
        for canon, spec in (cfg.get("brands") or {}).items():
            spec = spec or {}
            original = bool(spec.get("original", False))
            for a in [canon, *(spec.get("aliases") or [])]:
                self.alias[brand_text(a)] = (canon, original)
        self.unknown_exact = set()
        self.unknown_re = []
        for item in cfg.get("not_a_brand") or []:
            if str(item).startswith("^"):
                self.unknown_re.append(re.compile("^" + brand_text(item[1:])))
            else:
                self.unknown_exact.add(brand_text(item))

    def resolve(self, raw) -> tuple[str, str]:
        key = brand_text(raw)
        if not key or key in self.unknown_exact or any(r.search(key) for r in self.unknown_re):
            return "", self.UNKNOWN
        if key in self.alias:
            canon, original = self.alias[key]
            return canon, self.ORIGINAL if original else self.ANALOG
        return cell_text(raw), self.ANALOG


# Слова названия для подсказок и проверки «похожи ли названия».
STOP_WORDS = {"для", "под", "без", "или", "уаз", "газ", "газель", "компл", "комплект", "шт", "нов", "стар", "обр"}


def name_tokens(name: str) -> set[str]:
    words = re.findall(r"[a-zа-я]+|\d+", cell_text(name).lower().replace("ё", "е"))
    out = set()
    for w in words:
        if w.isdigit():
            if len(w) >= 2:
                out.add(w)
        elif len(w) >= 3 and w not in STOP_WORDS:
            out.add(w[:6])
    return out


# ---------------------------------------------------------------- чтение файлов


def excel_engine(path: Path) -> str:
    return {".xls": "xlrd", ".xlsb": "pyxlsb"}.get(path.suffix.lower(), "openpyxl")


def read_raw(path: Path, sheet) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, header=None, dtype=object, engine=excel_engine(path))


def find_header_row(raw: pd.DataFrame, names: list[str], limit: int = 60) -> int:
    wanted = {n.strip().lower() for n in names}
    for i in range(min(limit, len(raw))):
        cells = {cell_text(v).lower() for v in raw.iloc[i].tolist()}
        if wanted <= cells:
            return i
    raise ValueError(f"не найдена строка заголовков с колонками {sorted(names)}")


def table_from_raw(raw: pd.DataFrame, header_row, columns: dict, where: str) -> pd.DataFrame:
    """Сырой лист → таблица с колонками-ключами из columns и номером строки Excel."""
    names = [v for v in columns.values() if v]
    idx = find_header_row(raw, names) if header_row in (None, "auto") else int(header_row) - 1
    header = [cell_text(v) for v in raw.iloc[idx].tolist()]
    pos = {}
    for key, title in columns.items():
        if not title:
            continue
        hits = [i for i, h in enumerate(header) if h.lower() == str(title).strip().lower()]
        if not hits:
            raise ValueError(f"{where}: в строке {idx + 1} нет колонки «{title}». Есть: {[h for h in header if h]}")
        pos[key] = hits[0]
    body = raw.iloc[idx + 1 :]
    out = pd.DataFrame({key: body.iloc[:, i].tolist() for key, i in pos.items()})
    out["excel_row"] = [i + 1 for i in body.index]
    return out


@dataclass
class Offer:
    supplier: str
    sheet: str
    excel_row: int
    article_raw: str
    article: str
    keys: set[str]
    oem: list[str]
    brand_raw: str
    brand: str
    brand_kind: str
    name: str
    unit: str
    stock: float | None
    code: str
    price: float | None
    currency: str
    price_kzt: float | None
    pack: bool
    tokens: set[str] = field(default_factory=set)


def load_supplier(cfg: dict, settings: dict, brands: Brands) -> list[Offer]:
    path = resolve_file(cfg["file"])
    name = cfg["name"]
    currency = str(cfg.get("currency", settings["report_currency"])).upper()
    if currency == settings["report_currency"]:
        rate = 1.0
    elif currency in (settings.get("rates") or {}):
        rate = float(settings["rates"][currency])
    else:
        sys.exit(f"{name}: нет курса для валюты {currency} в settings.yaml → rates")
    columns = cfg["columns"]
    for required in ("article", "price"):
        if not columns.get(required):
            sys.exit(f"{name}: в columns не задана колонка {required}")

    sheets = cfg.get("sheets", "all")
    if sheets == "all":
        sheets = pd.ExcelFile(path, engine=excel_engine(path)).sheet_names
    pack_units = {str(u).strip().lower() for u in cfg.get("pack_units") or []}

    offers: list[Offer] = []
    seen = set()
    dedupe = cfg.get("dedupe_by")
    for spec in sheets:
        sheet = spec["name"] if isinstance(spec, dict) else spec
        header_row = spec.get("header_row", cfg.get("header_row")) if isinstance(spec, dict) else cfg.get("header_row")
        where = f"{name} / {path.name} / лист «{sheet}»"
        try:
            table = table_from_raw(read_raw(path, sheet), header_row, columns, where)
        except ValueError as e:
            if cfg.get("sheets", "all") == "all":
                print(f"  пропускаю лист: {e}")
                continue
            sys.exit(str(e))
        for row in table.to_dict("records"):
            article_raw = cell_text(row.get("article"))
            price = parse_price(row.get("price"))
            if not article_raw and price is None:
                continue  # заголовок группы, пустая строка
            if dedupe:
                key = cell_text(row.get(dedupe))
                if key and key in seen:
                    continue
                seen.add(key)
            brand, kind = brands.resolve(row.get("brand"))
            unit = cell_text(row.get("unit"))
            oem = [norm_article(x) for x in re.split(r"[;,|]", cell_text(row.get("oem"))) if norm_article(x)]
            offers.append(
                Offer(
                    supplier=name,
                    sheet=str(sheet),
                    excel_row=int(row["excel_row"]),
                    article_raw=article_raw,
                    article=norm_article(article_raw),
                    keys=factory_keys(article_raw),
                    oem=oem,
                    brand_raw=cell_text(row.get("brand")),
                    brand=brand,
                    brand_kind=kind,
                    name=cell_text(row.get("name")),
                    unit=unit,
                    stock=parse_number(row.get("stock")),
                    code=cell_text(row.get("code")),
                    price=price,
                    currency=currency,
                    price_kzt=round(price * rate, 2) if price is not None else None,
                    pack=unit.lower() in pack_units,
                )
            )
    print(f"  {name}: {path.name}, позиций {len(offers)}, с ценой {sum(o.price is not None for o in offers)}")
    return offers


@dataclass
class Product:
    id: str
    article_raw: str
    article: str
    keys: set[str]
    name: str
    car: str
    category: str
    price: float | None
    tokens: set[str]


def load_our(settings: dict) -> list[Product]:
    cfg = settings["our"]
    path = resolve_file(cfg["file"])
    table = table_from_raw(read_raw(path, cfg["sheet"]), cfg.get("header_row"), cfg["columns"], f"наш каталог / {path.name}")
    rate = 1.0 if cfg.get("currency", "KZT") == settings["report_currency"] else float(settings["rates"][cfg["currency"]])
    min_len = int(settings.get("min_article_length", 0))
    products = []
    for row in table.to_dict("records"):
        name = cell_text(row.get("name"))
        if not name:
            continue
        article_raw = cell_text(row.get("article"))
        article = norm_article(article_raw)
        if len(article) < min_len:
            article = ""  # слишком короткий: совпадёт с чем угодно
        price = parse_price(row.get("price"))
        products.append(
            Product(
                id=cell_text(row.get("id")),
                article_raw=article_raw,
                article=article,
                keys=factory_keys(article_raw) if article else set(),
                name=name,
                car=cell_text(row.get("car")),
                category=cell_text(row.get("category")),
                price=price * rate if price is not None else None,
                tokens=name_tokens(name),
            )
        )
    print(f"  наш каталог: {path.name}, товаров {len(products)}, с артикулом {sum(bool(p.article) for p in products)}")
    return products


# ---------------------------------------------------------------- сопоставление

M_EXACT = "артикул"
M_FACTORY = "артикул без точек, тире и кода модификации"
M_OEM = "кросс (OEM)"


@dataclass
class Match:
    product: Product
    offer: Offer
    method: str
    in_calc: bool = False
    notes: list[str] = field(default_factory=list)  # почему на проверку
    excluded: str = ""  # почему не в расчёте
    names_differ: bool = False


def match_all(products: list[Product], offers: list[Offer], settings: dict):
    by_article = defaultdict(list)
    by_key = defaultdict(list)
    by_oem = defaultdict(list)
    for i, o in enumerate(offers):
        if o.article:
            by_article[o.article].append(i)
        for k in o.keys:
            by_key[k].append(i)
        for x in o.oem:
            by_oem[x].append(i)

    doc_freq = Counter(t for o in offers for t in o.tokens)
    total = max(len(offers), 1)

    def idf(t: str) -> float:
        return math.log(total / (1 + doc_freq.get(t, 0))) + 1

    def similarity(p: Product, o: Offer) -> float:
        if not p.tokens:
            return 0.0
        weight = sum(idf(t) for t in p.tokens)
        return sum(idf(t) for t in p.tokens & o.tokens) / weight

    our_article_count = Counter(p.article for p in products if p.article)
    count_analogs = bool(settings.get("count_analogs", False))
    factor = float(settings.get("outlier_factor", 3))

    matches: list[Match] = []
    for p in products:
        found: dict[int, str] = {}
        if p.article:
            for i in by_article.get(p.article, []):
                found.setdefault(i, M_EXACT)
            for k in p.keys:
                for i in by_key.get(k, []):
                    found.setdefault(i, M_FACTORY)
            for i in by_oem.get(p.article, []):
                found.setdefault(i, M_OEM)
        own = []
        for i, method in found.items():
            o = offers[i]
            m = Match(p, o, method)
            if o.price is None:
                m.excluded = "нет цены в прайсе"
            elif o.brand_kind == Brands.UNKNOWN:
                m.excluded = "производитель не указан или указан не бренд"
                m.notes.append(f"производитель в прайсе: «{o.brand_raw or 'пусто'}»")
            elif o.brand_kind == Brands.ANALOG and not count_analogs:
                m.excluded = "аналог, не ГАЗ/УАЗ"
            else:
                m.in_calc = True
            if o.pack:
                m.notes.append(f"цена за «{o.unit}», возможно не за штуку")
            if our_article_count[p.article] > 1:
                m.notes.append("этот артикул у нас у нескольких товаров")
            if o.tokens and p.tokens and similarity(p, o) == 0:
                m.names_differ = True
                m.notes.append("в названиях нет общих слов")
            own.append(m)

        # Выбросы: дальше от медианы, чем в factor раз. При двух ценах
        # не понять, какая из них ошибочная, — обе остаются и идут на проверку.
        calc = [m for m in own if m.in_calc]
        if len(calc) >= 2:
            med = median(m.offer.price_kzt for m in calc)
            far = f"отличается от медианы {med:,.0f} ₸ больше чем в {factor:g} раза".replace(",", " ")
            for m in calc:
                if m.offer.price_kzt > med * factor or m.offer.price_kzt < med / factor:
                    if len(calc) >= 3:
                        m.in_calc = False
                        m.excluded = f"выброс: {far}"
                        m.notes.append("вероятная ошибка цены")
                    else:
                        m.notes.append(f"цены расходятся: {far}, какая ошибочная — не понять")
            if len(calc) == 2 and any("цены расходятся" in n for m in calc for n in m.notes):
                for m in calc:
                    if not any("цены расходятся" in n for n in m.notes):
                        m.notes.append(f"цены расходятся: {far}, какая ошибочная — не понять")
        matches.extend(own)

    return matches, similarity


def name_hints(products, offers, matched_ids, similarity, settings) -> dict[str, list[tuple[float, Offer]]]:
    limit = int(settings.get("name_hints", 3))
    threshold = float(settings.get("name_hint_min_score", 0.5))
    by_token = defaultdict(list)
    for i, o in enumerate(offers):
        if o.price is not None:
            for t in o.tokens:
                by_token[t].append(i)
    hints = {}
    for p in products:
        if p.id in matched_ids or not p.tokens:
            continue
        candidates = set()
        for t in p.tokens:
            if len(by_token[t]) < 3000:  # слишком частое слово не сужает поиск
                candidates.update(by_token[t])
        scored = sorted(
            ((similarity(p, offers[i]), offers[i]) for i in candidates),
            key=lambda x: (-x[0], x[1].brand_kind != Brands.ORIGINAL, len(x[1].name)),
        )
        hints[p.id] = [(s, o) for s, o in scored if s >= threshold][:limit]
    return hints


# ---------------------------------------------------------------- расчёт


HAS_SEPARATOR = re.compile(r"[\s\-.]")


def supplier_article(own: list[Match]) -> tuple[str, str]:
    """Как артикул пишут поставщики: с точками и тире, по которым его находит поиск.

    Берётся из совпадений с похожим названием, сначала у ГАЗ/УАЗ.
    Возвращает (артикул, пример: поставщик и наименование).
    """
    for original_only in (True, False):
        pool = [
            m for m in own
            if not m.names_differ
            and HAS_SEPARATOR.search(m.offer.article_raw)
            and (m.offer.brand_kind == Brands.ORIGINAL or not original_only)
        ]
        if pool:
            spelled = Counter(re.sub(r"\s+", " ", m.offer.article_raw) for m in pool)
            best = spelled.most_common(1)[0][0]
            example = next(m for m in pool if re.sub(r"\s+", " ", m.offer.article_raw) == best)
            return best, f"{example.offer.supplier}: {example.offer.name}"
    return "", ""


def summarize(products, matches, suppliers):
    by_product = defaultdict(list)
    for m in matches:
        by_product[m.product.id].append(m)
    rows = []
    for p in products:
        own = by_product.get(p.id, [])
        per_supplier = {}
        for m in own:
            if m.in_calc:
                cur = per_supplier.get(m.offer.supplier)
                if cur is None or m.offer.price_kzt < cur:
                    per_supplier[m.offer.supplier] = m.offer.price_kzt
        analogs = [m.offer.price_kzt for m in own if m.offer.brand_kind == Brands.ANALOG and m.offer.price_kzt]
        prices = list(per_supplier.values())
        r = {
            "p": p,
            "n": len(per_supplier),
            "per_supplier": per_supplier,
            "analogs": len(analogs),
            "analog_min": min(analogs) if analogs else None,
            "review": sum(bool(m.notes) for m in own),
            "supplier_article": supplier_article(own),
        }
        if prices:
            lo = min(prices)
            r.update(
                min=lo,
                max=max(prices),
                mean=mean(prices),
                median=median(prices),
                who_min=", ".join(s for s in suppliers if per_supplier.get(s) == lo),
            )
            if p.price:
                r["dev_mean"] = p.price / r["mean"] - 1
                r["dev_median"] = p.price / r["median"] - 1
            r["status"] = "сопоставлено"
        elif any(m.offer.brand_kind == Brands.ANALOG and m.offer.price is not None for m in own) and not any(
            m.offer.brand_kind != Brands.ANALOG and m.offer.price is not None for m in own
        ):
            r["status"] = "только аналоги"
        elif any(m.offer.price is not None for m in own):
            r["status"] = "на проверку"
        else:
            r["status"] = "не найдено"
        rows.append(r)
    return rows


# ---------------------------------------------------------------- Excel

HEAD_FILL = PatternFill("solid", fgColor="DDE3EA")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6EFCE")
MONEY = "#,##0"
PCT = "+0%;-0%;0%"


def write_sheet(ws, headers, rows, widths=None, formats=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row, 1)}"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) + 2)))
        fmt = (formats or {}).get(h)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt


def detail_row(m: Match) -> list:
    o, p = m.offer, m.product
    return [
        p.id, p.article_raw, p.name, p.price,
        o.supplier, o.article_raw, o.brand_raw, o.brand, o.brand_kind, o.name, o.unit, o.stock,
        o.price, o.currency, o.price_kzt, m.method,
        "да" if m.in_calc else "нет", m.excluded, "; ".join(m.notes),
        f"{o.sheet}, строка {o.excel_row}",
    ]


DETAIL_HEADERS = [
    "ID", "Наш артикул", "Наше наименование", "Наша цена, ₸",
    "Поставщик", "Артикул поставщика", "Производитель (в прайсе)", "Бренд", "Тип бренда",
    "Наименование у поставщика", "Ед.", "Остаток",
    "Цена в прайсе", "Валюта", "Цена, ₸", "Способ совпадения",
    "В расчёте", "Почему не в расчёте", "На проверку: причина", "Где в прайсе",
]
DETAIL_WIDTHS = {
    "Наше наименование": 34, "Наименование у поставщика": 50, "Производитель (в прайсе)": 22,
    "Почему не в расчёте": 34, "На проверку: причина": 44, "Способ совпадения": 24, "Где в прайсе": 22,
}
DETAIL_FORMATS = {"Наша цена, ₸": MONEY, "Цена, ₸": MONEY, "Цена в прайсе": "#,##0.00"}


def markup_stats(pairs: list[tuple[float, float]]) -> list:
    """Наценка нашей цены к цене поставщика: [товаров, медиана, средняя, мы дороже, мы дешевле]."""
    pct = [ours / theirs - 1 for ours, theirs in pairs if ours and theirs]
    if not pct:
        return [0, None, None, 0, 0]
    return [len(pct), median(pct), mean(pct), sum(x > 0 for x in pct), sum(x < 0 for x in pct)]


def write_totals(ws, rows, suppliers, settings):
    """Насколько наша цена выше оптовой: в целом и по каждому поставщику."""
    matched = [r for r in rows if r["n"]]
    clean = [r for r in matched if not r["review"]]
    body = [
        ["Средняя цена поставщиков", *markup_stats([(r["p"].price, r["mean"]) for r in matched])],
        ["Средняя цена поставщиков, без строк на проверку", *markup_stats([(r["p"].price, r["mean"]) for r in clean])],
        ["Минимальная цена поставщиков", *markup_stats([(r["p"].price, r["min"]) for r in matched])],
        *[[s, *markup_stats([(r["p"].price, r["per_supplier"][s]) for r in matched if s in r["per_supplier"]])]
          for s in suppliers],
    ]
    headers = ["Сравниваем с", "Товаров", "Наценка, медиана", "Наценка, средняя", "Мы дороже", "Мы дешевле"]
    write_sheet(ws, headers, body, {"Сравниваем с": 48, "Наценка, медиана": 18, "Наценка, средняя": 18},
                {"Наценка, медиана": PCT, "Наценка, средняя": PCT})
    ws.auto_filter.ref = None
    notes = [
        "Наценка = наша цена / цена поставщика − 1. +40% значит, что мы на 40% дороже.",
        f"Наша цена: {settings['our'].get('price_note', '')}. Цены поставщиков: {settings.get('supplier_price_note', '')}.",
        "Медиана надёжнее средней: одна ошибочная цена её почти не сдвигает.",
        "Считаются только товары, найденные у поставщиков с производителем ГАЗ/УАЗ.",
    ]
    ws.append([])
    for n in notes:
        ws.append([n])


def write_report(path, products, matches, rows, hints, suppliers, settings, sources):
    wb = Workbook()

    ws = wb.active
    ws.title = "Итоги"
    write_totals(ws, rows, suppliers, settings)

    ws = wb.create_sheet("Сводка")
    sup_cols = [f"{s}, ₸" for s in suppliers]
    headers = [
        "ID", "Артикул", "Артикул у поставщиков", "Наименование", "Подходит к авто", "Категория", "Наша цена, ₸",
        "Найдено у поставщиков", "Мин, ₸", "У кого минимум", "Макс, ₸", "Средняя, ₸", "Медиана, ₸",
        "Наценка к средней", "Наценка к медиане", *sup_cols,
        "Аналогов (не в расчёте)", "Мин. цена аналога, ₸", "Строк на проверку", "Статус",
    ]
    body = []
    for r in rows:
        p = r["p"]
        body.append([
            p.id, p.article_raw, r["supplier_article"][0] or None, p.name, p.car, p.category, p.price,
            r["n"], r.get("min"), r.get("who_min"), r.get("max"), r.get("mean"), r.get("median"),
            r.get("dev_mean"), r.get("dev_median"), *[r["per_supplier"].get(s) for s in suppliers],
            r["analogs"] or None, r["analog_min"], r["review"] or None, r["status"],
        ])
    fmts = {h: MONEY for h in ["Наша цена, ₸", "Мин, ₸", "Макс, ₸", "Средняя, ₸", "Медиана, ₸", "Мин. цена аналога, ₸", *sup_cols]}
    fmts.update({"Наценка к средней": PCT, "Наценка к медиане": PCT})
    write_sheet(ws, headers, body, {"Наименование": 40, "Артикул": 18, "Артикул у поставщиков": 22, "У кого минимум": 18, "Подходит к авто": 14}, fmts)
    width = len(headers)
    for i, r in enumerate(rows, start=2):
        dev = r.get("dev_mean")
        if dev is None or abs(dev) < 1e-9:
            continue
        fill = RED if dev > 0 else GREEN
        for col in range(1, width + 1):
            ws.cell(row=i, column=col).fill = fill

    ws = wb.create_sheet("Детали")
    write_sheet(ws, DETAIL_HEADERS, [detail_row(m) for m in matches], DETAIL_WIDTHS, DETAIL_FORMATS)

    ws = wb.create_sheet("На проверку")
    review = [m for m in matches if m.notes]
    write_sheet(ws, DETAIL_HEADERS, [detail_row(m) for m in review], DETAIL_WIDTHS, DETAIL_FORMATS)

    ws = wb.create_sheet("Не найдено")
    limit = int(settings.get("name_hints", 3))
    nf_headers = ["ID", "Артикул", "Наименование", "Подходит к авто", "Категория", "Наша цена, ₸", "Причина",
                  *[f"Подсказка по названию {i}" for i in range(1, limit + 1)]]
    nf_rows = []
    for r in rows:
        if r["status"] != "не найдено":
            continue
        p = r["p"]
        reason = "артикул не найден в прайсах" if p.article else ("нет артикула" if not p.article_raw else "артикул слишком короткий")
        hint_cells = [
            f"{o.supplier}: {o.name} [{o.article_raw or 'без артикула'}; {o.brand_raw or 'производитель не указан'}] "
            f"— {o.price_kzt:,.0f} ₸ (сходство {s:.0%})".replace(",", " ")
            for s, o in hints.get(p.id, [])
        ]
        nf_rows.append([p.id, p.article_raw, p.name, p.car, p.category, p.price, reason,
                        *hint_cells, *[None] * (limit - len(hint_cells))])
    write_sheet(ws, nf_headers, nf_rows,
                {"Наименование": 34, "Причина": 26, **{f"Подсказка по названию {i}": 70 for i in range(1, limit + 1)}},
                {"Наша цена, ₸": MONEY})

    ws = wb.create_sheet("Артикулы")
    art_rows = [
        [r["p"].id, r["p"].article_raw, r["supplier_article"][0], r["p"].name, r["supplier_article"][1]]
        for r in rows
        if r["supplier_article"][0] and not HAS_SEPARATOR.search(r["p"].article_raw)
    ]
    write_sheet(
        ws,
        ["ID", "Наш артикул", "Как пишут поставщики", "Наше наименование", "Пример у поставщика"],
        art_rows,
        {"Наш артикул": 18, "Как пишут поставщики": 24, "Наше наименование": 40, "Пример у поставщика": 80},
    )

    ws = wb.create_sheet("Параметры")
    info = [
        ["Дата расчёта", date.today().isoformat()],
        ["Валюта отчёта", settings["report_currency"]],
        *[[f"Курс {c}", f"{v} ₸ ({settings.get('rates_source', '')}, {settings.get('rates_date', '')})"]
          for c, v in (settings.get("rates") or {}).items()],
        ["Наша цена", settings["our"].get("price_note", "")],
        ["Цены поставщиков", settings.get("supplier_price_note", "")],
        ["Аналоги в расчёте", "да" if settings.get("count_analogs") else "нет, только ГАЗ/УАЗ"],
        ["Выброс", f"цена дальше от медианы больше чем в {settings.get('outlier_factor', 3)} раза"],
        *[[f"Файл: {k}", v] for k, v in sources.items()],
        ["Подсказки по названию", "в расчёт не идут"],
    ]
    write_sheet(ws, ["Параметр", "Значение"], info, {"Параметр": 34, "Значение": 90})
    ws.auto_filter.ref = None

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- запуск


def main() -> None:
    settings = load_yaml("settings.yaml")
    brands = Brands(load_yaml("brands.yaml"))
    supplier_cfgs = load_yaml("suppliers.yaml").get("suppliers") or []
    if not supplier_cfgs:
        sys.exit("В config/suppliers.yaml нет ни одного поставщика")

    print("Читаю файлы")
    products = load_our(settings)
    offers: list[Offer] = []
    sources = {"наш каталог": settings["our"]["file"]}
    for cfg in supplier_cfgs:
        loaded = load_supplier(cfg, settings, brands)
        offers.extend(loaded)
        sources[cfg["name"]] = str(resolve_file(cfg["file"]).relative_to(ROOT))
    for o in offers:
        o.tokens = name_tokens(o.name)
    suppliers = [c["name"] for c in supplier_cfgs]

    print("Сопоставляю")
    matches, similarity = match_all(products, offers, settings)
    matched_ids = {m.product.id for m in matches}
    hints = name_hints(products, offers, matched_ids, similarity, settings)
    rows = summarize(products, matches, suppliers)

    out = ROOT / settings.get("output_dir", "result") / f"price_comparison_{date.today().isoformat()}.xlsx"
    write_report(out, products, matches, rows, hints, suppliers, settings, sources)

    status = Counter(r["status"] for r in rows)
    flagged = sum(1 for r in rows if r["status"] == "сопоставлено" and r["review"])
    print()
    print(f"Товаров в каталоге: {len(products)}")
    print(f"  сопоставлено (цена в расчёте): {status['сопоставлено']}, из них с пометками на проверку: {flagged}")
    print(f"  только на проверку (совпадения есть, в расчёт не вошли): {status['на проверку']}")
    print(f"  только аналоги других производителей: {status['только аналоги']}")
    print(f"  не найдено: {status['не найдено']}")
    for s in suppliers:
        print(f"  найдено у «{s}»: {sum(1 for r in rows if s in r['per_supplier'])}")
    print(f"Строк совпадений: {len(matches)}, на проверку: {sum(bool(m.notes) for m in matches)}")
    print(f"Результат: {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
